#!/usr/bin/env python3
"""
MaaS 消费明细解析 & HTML 报告生成工具
- 按 Payer user name / Resource ID / Product / Instance ID 等多维度分类统计
- 输出自包含 HTML 报告（含表格 + Canvas 柱状图）
"""

import csv
import os
import re
import sys
from collections import defaultdict
from datetime import datetime

# ── 通用统计结构 ─────────────────────────────────────────────────────
def make_stat():
    return {
        "count": 0,
        "total_usage": 0.0,
        "total_pretax": 0.0,
        "total_tax": 0.0,
        "total_posttax": 0.0,
        "total_amount_due": 0.0,
        "total_amount_paid": 0.0,
        "usage_unit": None,
        "currency": None,
        "display_name": "",
    }


def make_sub_stat():
    """子资源统计结构"""
    return {
        "count": 0,
        "total_usage": 0.0,
        "total_amount_due": 0.0,
        "total_pretax": 0.0,
        "usage_unit": None,
        "currency": None,
    }


def _detect_delimiter(first_line: str) -> str:
    """自动检测分隔符：统计第一行中 tab 和逗号的数量，取较多的那个"""
    tabs = first_line.count("\t")
    commas = first_line.count(",")
    return "\t" if tabs >= commas else ","


def load_mapping(filepath: str):
    """加载名称→Resource ID 映射表，返回 (mapping, resource_to_username)
       mapping: {resource_id: name}
       resource_to_username: {resource_id: username}  从 api-key-{username}-{N} 中提取 username"""
    mapping = {}
    resource_to_username = {}
    script_dir = os.path.dirname(os.path.abspath(__file__))
    full_path = os.path.join(script_dir, filepath)
    if not os.path.exists(full_path):
        print(f"[WARN] 映射表未找到: {full_path}")
        return mapping, resource_to_username
    with open(full_path, "r", encoding="utf-8-sig") as f:
        first_line = f.readline()
        delimiter = _detect_delimiter(first_line)
        f.seek(0)
        reader = csv.DictReader(f, delimiter=delimiter)
        for r in reader:
            keys = list(r.keys())
            name = (r.get("名称") or (keys[0] if len(keys) > 0 else "") or "").strip()
            rid = (r.get("资源ID") or (keys[1] if len(keys) > 1 else "") or "").strip()
            if rid and name:
                mapping[rid] = name
                # 提取 username: api-key-日本电信-5 → 日本电信
                username = extract_username(name)
                if username:
                    resource_to_username[rid] = username
    print(f"  加载映射表: {len(mapping)} 条 (分隔符: {'Tab' if delimiter == chr(9) else '逗号'}), 提取用户名: {len(set(resource_to_username.values()))} 个")
    return mapping, resource_to_username


def extract_username(name: str) -> str:
    """从映射名称中提取用户名: api-key-日本电信-5 → 日本电信, api-key-flatkey → flatkey"""
    # 去掉 api-key- 前缀
    if name.startswith("api-key-"):
        rest = name[len("api-key-"):]
    else:
        rest = name
    # 去掉末尾的 -数字（如 -5, -12）
    rest = re.sub(r'-\d+$', '', rest)
    return rest


def safe_float(s):
    try:
        return float(s) if s and s != "-" else 0.0
    except ValueError:
        return 0.0


# ── USAGE 单位标准化 ─────────────────────────────────────────────────
def normalize_usage(raw):
    """返回 (数值, 单位)"""
    if raw in (None, "-", ""):
        return (0.0, "")
    raw = raw.strip()
    # 尝试纯数字或已带单位的格式（本文件 Usage 列是纯数字，单位在 Usage unit 列）
    for ch in ",，" :
        raw = raw.replace(ch, "")
    return (safe_float(raw), "")


# ── 解析 CSV ─────────────────────────────────────────────────────────
def parse_csv(filepath: str, resource_to_username: dict = None):
    if resource_to_username is None:
        resource_to_username = {}
    rows = []
    with open(filepath, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            usage_val, _ = normalize_usage(r.get("Usage", ""))
            resource_id = (r.get("Resource ID") or "").strip()
            user_name = resource_to_username.get(resource_id, "")
            rows.append(
                {
                    "billing_cycle": (r.get("Billing cycle") or "").strip(),
                    "payer_user_name": (r.get("Payer user name") or "").strip(),
                    "payer_org": (r.get("Payer organization name") or "").strip(),
                    "product": (r.get("Product") or "").strip(),
                    "business_mode": (r.get("Business mode") or "").strip(),
                    "start_time": (r.get("Consumption start time(UTC+8)") or "").strip(),
                    "end_time": (r.get("Consumption end time(UTC+8)") or "").strip(),
                    "instance_id": (r.get("Instance ID") or "").strip(),
                    "instance_name": (r.get("Instance name") or "").strip(),
                    "resource_id": resource_id,
                    "user_name": user_name,
                    "split_item_id": (r.get("Split item id") or "").strip(),
                    "split_item_name": (r.get("Split item name") or "").strip(),
                    "usage": usage_val,
                    "usage_unit": (r.get("Usage unit") or "").strip(),
                    "usage_pct": safe_float(r.get("Percentage of split item usage")),
                    "exchange_rate": safe_float(r.get("Exchange rate")),
                    "currency": (r.get("Payment Currency") or "").strip(),
                    "pretax_amount": safe_float(r.get("Pre-tax amount (payment currency)")),
                    "pretax_real": safe_float(r.get("Pre-tax real value (payment currency)")),
                    "tax": safe_float(r.get("Tax (payment currency)")),
                    "posttax_real": safe_float(r.get("Post-tax real value (payment currency)")),
                    "posttax_amount": safe_float(r.get("Post-tax amount (payment currency)")),
                    "amount_paid": safe_float(r.get("Amount paid")),
                    "credit_refund": safe_float(r.get("Credit Refund Deduction")),
                    "amount_due": safe_float(r.get("Amount due")),
                    "project": (r.get("Project") or "").strip(),
                    "project_display": (r.get("Project display name") or "").strip(),
                    "tag": (r.get("Tag") or "").strip(),
                }
            )
    return rows


def parse_xlsx(filepath: str, resource_to_username: dict = None):
    """解析 .xlsx 文件，与 parse_csv 返回相同结构"""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("解析 xlsx 需要 openpyxl，请执行: pip install openpyxl")
    if resource_to_username is None:
        resource_to_username = {}

    def _s(v):
        """安全转字符串，处理 openpyxl 返回的 int/float/datetime"""
        if v is None:
            return ""
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return str(v)

    rows = []
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    # 读取表头
    headers = []
    for cell in ws[1]:
        headers.append((cell.value or "").strip())
    # 解析数据行
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(headers, row))
        usage_val, _ = normalize_usage(_s(r.get("Usage")))
        resource_id = _s(r.get("Resource ID")).strip()
        user_name = resource_to_username.get(resource_id, "")
        rows.append(
            {
                "billing_cycle": _s(r.get("Billing cycle")).strip(),
                "payer_user_name": _s(r.get("Payer user name")).strip(),
                "payer_org": _s(r.get("Payer organization name")).strip(),
                "product": _s(r.get("Product")).strip(),
                "business_mode": _s(r.get("Business mode")).strip(),
                "start_time": _s(r.get("Consumption start time(UTC+8)")).strip(),
                "end_time": _s(r.get("Consumption end time(UTC+8)")).strip(),
                "instance_id": _s(r.get("Instance ID")).strip(),
                "instance_name": _s(r.get("Instance name")).strip(),
                "resource_id": resource_id,
                "user_name": user_name,
                "split_item_id": _s(r.get("Split item id")).strip(),
                "split_item_name": _s(r.get("Split item name")).strip(),
                "usage": usage_val,
                "usage_unit": _s(r.get("Usage unit")).strip(),
                "usage_pct": safe_float(r.get("Percentage of split item usage")),
                "exchange_rate": safe_float(r.get("Exchange rate")),
                "currency": _s(r.get("Payment Currency")).strip(),
                "pretax_amount": safe_float(r.get("Pre-tax amount (payment currency)")),
                "pretax_real": safe_float(r.get("Pre-tax real value (payment currency)")),
                "tax": safe_float(r.get("Tax (payment currency)")),
                "posttax_real": safe_float(r.get("Post-tax real value (payment currency)")),
                "posttax_amount": safe_float(r.get("Post-tax amount (payment currency)")),
                "amount_paid": safe_float(r.get("Amount paid")),
                "credit_refund": safe_float(r.get("Credit Refund Deduction")),
                "amount_due": safe_float(r.get("Amount due")),
                "project": _s(r.get("Project")).strip(),
                "project_display": _s(r.get("Project display name")).strip(),
                "tag": _s(r.get("Tag")).strip(),
            }
        )
    wb.close()
    return rows


def parse_file(filepath: str, resource_to_username: dict = None):
    """统一入口：根据扩展名自动选择 CSV 或 XLSX 解析"""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".xlsx":
        return parse_xlsx(filepath, resource_to_username)
    else:
        return parse_csv(filepath, resource_to_username)


# ── 聚合函数 ─────────────────────────────────────────────────────────
def aggregate(rows: list, key_field: str, top_n: int = 30, mapping: dict = None):
    if mapping is None:
        mapping = {}
    groups = defaultdict(make_stat)
    # user_name 维度额外追踪子 resource_id 明细
    sub_groups = defaultdict(lambda: defaultdict(make_sub_stat))

    for row in rows:
        key = row[key_field] or "(空)"
        s = groups[key]
        s["count"] += 1
        s["total_usage"] += row["usage"]
        s["total_pretax"] += row["pretax_amount"]
        s["total_tax"] += row["tax"]
        s["total_posttax"] += row["posttax_amount"]
        s["total_amount_due"] += row["amount_due"]
        s["total_amount_paid"] += row["amount_paid"]
        if s["usage_unit"] is None and row["usage_unit"]:
            s["usage_unit"] = row["usage_unit"]
        if s["currency"] is None and row["currency"]:
            s["currency"] = row["currency"]
        if key_field == "resource_id" and key in mapping:
            s["display_name"] = mapping[key]

        # user_name 维度：记录子 resource_id 明细
        if key_field == "user_name" and row["resource_id"]:
            rid = row["resource_id"]
            ss = sub_groups[key][rid]
            ss["count"] += 1
            ss["total_usage"] += row["usage"]
            ss["total_amount_due"] += row["amount_due"]
            ss["total_pretax"] += row["pretax_amount"]
            if ss["usage_unit"] is None and row["usage_unit"]:
                ss["usage_unit"] = row["usage_unit"]
            if ss["currency"] is None and row["currency"]:
                ss["currency"] = row["currency"]

    result = [
        {
            "key": k,
            "display_name": v["display_name"],
            "count": v["count"],
            "total_usage": round(v["total_usage"], 2),
            "total_pretax": round(v["total_pretax"], 2),
            "total_tax": round(v["total_tax"], 2),
            "total_posttax": round(v["total_posttax"], 2),
            "total_amount_due": round(v["total_amount_due"], 2),
            "total_amount_paid": round(v["total_amount_paid"], 2),
            "avg_usage": round(v["total_usage"] / v["count"], 2) if v["count"] else 0,
            "usage_unit": v["usage_unit"] or "",
            "currency": v["currency"] or "",
            "sub_resources": [
                {
                    "resource_id": rid,
                    "display_name": mapping.get(rid, ""),
                    "count": ss["count"],
                    "total_usage": round(ss["total_usage"], 2),
                    "total_amount_due": round(ss["total_amount_due"], 2),
                    "total_pretax": round(ss["total_pretax"], 2),
                    "usage_unit": ss["usage_unit"] or "",
                    "currency": ss["currency"] or "",
                }
                for rid, ss in sorted(sub_groups[k].items(), key=lambda x: x[1]["total_usage"], reverse=True)
            ] if key_field == "user_name" else [],
        }
        for k, v in groups.items()
    ]
    return sorted(result, key=lambda x: x["total_amount_due"], reverse=True)[:top_n]


# ── HTML 生成 ─────────────────────────────────────────────────────────
def build_html(rows, stats, csv_filename):
    total_records = len(rows)
    total_usage = round(sum(r["usage"] for r in rows), 2)
    total_amount_due = round(sum(r["amount_due"] for r in rows), 2)
    total_pretax = round(sum(r["pretax_amount"] for r in rows), 2)
    total_posttax = round(sum(r["posttax_amount"] for r in rows), 2)
    total_tax = round(sum(r["tax"] for r in rows), 2)
    total_amount_paid = round(sum(r["amount_paid"] for r in rows), 2)
    avg_usage = round(total_usage / total_records, 2) if total_records else 0
    unique_users = len(set(r["payer_user_name"] for r in rows))
    unique_products = len(set(r["product"] for r in rows))
    unique_resources = len(set(r["resource_id"] for r in rows))
    unique_instances = len(set(r["instance_id"] for r in rows))
    unique_projects = len(set(r["project"] for r in rows if r["project"] and r["project"] != "-"))

    # 时间范围
    times = [r["start_time"] for r in rows if r["start_time"]] + \
            [r["end_time"] for r in rows if r["end_time"]]
    time_range = ""
    if times:
        time_range = f"{min(times)} ~ {max(times)}"

    # 获取主要币种和 Usage 单位
    main_currency = stats["payer_user_name"][0]["currency"] if stats["payer_user_name"] else ""
    main_unit = stats["payer_user_name"][0]["usage_unit"] if stats["payer_user_name"] else ""

    # ── 表格行 HTML ──
    def table_rows(data, show_currency=True):
        lines = []
        for i, d in enumerate(data):
            currency = f" {d['currency']}" if show_currency and d.get("currency") else ""
            display = d.get("display_name") or d["key"]
            title = d["key"] if d.get("display_name") else display
            lines.append(
                f"""<tr>
                <td>{i + 1}</td>
                <td class="key-cell" title="{title}">{display}</td>
                <td class="num">{d['count']:,}</td>
                <td class="num">{d['total_usage']:,} {d['usage_unit']}</td>
                <td class="num">{d['avg_usage']:,} {d['usage_unit']}</td>
                <td class="num">{d['total_pretax']:,.2f}{currency}</td>
                <td class="num">{d['total_tax']:,.2f}{currency}</td>
                <td class="num">{d['total_posttax']:,.2f}{currency}</td>
                <td class="num">{d['total_amount_due']:,.2f}{currency}</td>
            </tr>"""
            )
        return "\n".join(lines)

    # ── 可展开表格行 HTML（user-username 维度） ──
    def user_table_rows(data):
        lines = []
        for i, d in enumerate(data):
            currency = f" {d['currency']}" if d.get("currency") else ""
            display = d.get("display_name") or d["key"]
            title = d["key"] if d.get("display_name") else display
            sub_count = len(d.get("sub_resources", []))
            row_id = f"user-row-{i}"
            lines.append(
                f"""<tr class="expandable-row" data-target="{row_id}" onclick="toggleSub('{row_id}',this)">
                <td class="toggle-icon">&#9654;</td>
                <td>{i + 1}</td>
                <td class="key-cell" title="{title}">{display} <span class="sub-badge">{sub_count} keys</span></td>
                <td class="num">{d['count']:,}</td>
                <td class="num">{d['total_usage']:,} {d['usage_unit']}</td>
                <td class="num">{d['avg_usage']:,} {d['usage_unit']}</td>
                <td class="num">{d['total_pretax']:,.2f}{currency}</td>
                <td class="num">{d['total_tax']:,.2f}{currency}</td>
                <td class="num">{d['total_posttax']:,.2f}{currency}</td>
                <td class="num">{d['total_amount_due']:,.2f}{currency}</td>
            </tr>"""
            )
            # 子行：resource_id 明细
            for j, sr in enumerate(d.get("sub_resources", [])):
                sr_currency = f" {sr['currency']}" if sr.get("currency") else ""
                sr_display = sr.get("display_name") or sr["resource_id"]
                sr_title = sr["resource_id"] if sr.get("display_name") else sr_display
                lines.append(
                    f"""<tr class="sub-row hidden" data-parent="{row_id}">
                <td></td>
                <td></td>
                <td class="key-cell sub-key" title="{sr_title}">&nbsp;&nbsp;&#9492; {sr_display}</td>
                <td class="num sub-num">{sr['count']:,}</td>
                <td class="num sub-num">{sr['total_usage']:,} {sr['usage_unit']}</td>
                <td class="num sub-num"></td>
                <td class="num sub-num">{sr['total_pretax']:,.2f}{sr_currency}</td>
                <td class="num sub-num"></td>
                <td class="num sub-num"></td>
                <td class="num sub-num">{sr['total_amount_due']:,.2f}{sr_currency}</td>
            </tr>"""
                )
        return "\n".join(lines)

    # ── 柱状图 Canvas 数据 ──
    def chart_data_json(data, chart_id, value_field, label, top=10):
        top_data = data[:top]
        labels = [d.get("display_name") or d["key"] for d in top_data]
        values = [d[value_field] for d in top_data]
        return f"""{{ id: "{chart_id}", labels: {labels}, values: {values}, label: "{label}" }}"""

    payer_chart = chart_data_json(stats["payer_user_name"], "chartPayer", "total_amount_due", "Amount Due", 10)
    resource_chart = chart_data_json(stats["resource_id"], "chartResource", "total_usage", "Total Usage", 10)
    product_chart = chart_data_json(stats["product"], "chartProduct", "total_amount_due", "Amount Due", 10)

    html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>MaaS 消费明细分析报告</title>
<style>
  :root {{
    --bg: #f5f7fa;
    --card: #fff;
    --text: #333;
    --muted: #888;
    --accent: #4f6ef7;
    --accent2: #22c55e;
    --accent3: #f59e0b;
    --border: #e5e7eb;
    --table-stripe: #f9fafb;
  }}
  * {{margin:0;padding:0;box-sizing:border-box;}}
  body {{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif;background:var(--bg);color:var(--text);line-height:1.6;}}
  .container {{max-width:1400px;margin:0 auto;padding:20px;}}
  h1 {{font-size:26px;margin:20px 0 8px;}}
  .subtitle {{color:var(--muted);font-size:14px;margin-bottom:24px;}}

  /* 概览卡片 */
  .cards {{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:14px;margin-bottom:30px;}}
  .card {{background:var(--card);border-radius:12px;padding:16px 18px;box-shadow:0 1px 3px rgba(0,0,0,.06);border:1px solid var(--border);}}
  .card-wide {{grid-column:span 2;}}
  .card .label {{font-size:11px;color:var(--muted);text-transform:uppercase;letter-spacing:.5px;line-height:1.3;}}
  .card .value {{font-size:26px;font-weight:700;margin-top:6px;}}
  .card .value.accent {{color:var(--accent);}}
  .card .value.accent2 {{color:var(--accent2);}}
  .card .value.accent3 {{color:var(--accent3);}}
  .card .value.time-val {{font-size:14px;font-weight:600;word-break:break-all;}}

  /* 图表区域 */
  .charts {{display:grid;grid-template-columns:1fr;gap:20px;margin-bottom:30px;}}
  .chart-box {{background:var(--card);border-radius:12px;padding:20px 16px;box-shadow:0 1px 3px rgba(0,0,0,.06);border:1px solid var(--border);overflow:hidden;}}
  .chart-box h3 {{font-size:15px;margin-bottom:12px;color:var(--muted);}}
  canvas {{width:100%;height:300px;display:block;}}

  /* 表格区域 */
  .section {{margin-bottom:36px;}}
  .section h2 {{font-size:18px;margin-bottom:12px;display:flex;align-items:center;gap:8px;}}
  .section h2 .badge {{font-size:12px;background:var(--accent);color:#fff;padding:2px 10px;border-radius:99px;}}
  .table-wrap {{overflow-x:auto;background:var(--card);border-radius:12px;box-shadow:0 1px 3px rgba(0,0,0,.06);border:1px solid var(--border);}}
  table {{width:100%;border-collapse:collapse;font-size:13px;}}
  th {{background:#f1f5f9;padding:10px 12px;text-align:left;font-weight:600;white-space:nowrap;color:var(--muted);font-size:12px;text-transform:uppercase;}}
  th.num {{text-align:right;}}
  td {{padding:8px 12px;border-top:1px solid var(--border);white-space:nowrap;}}
  tr:hover {{background:#eef2ff;}}
  .key-cell {{max-width:260px;overflow:hidden;text-overflow:ellipsis;font-family:monospace;font-size:12px;}}
  .num {{text-align:right;font-variant-numeric:tabular-nums;}}

  /* 可展开行 */
  .expandable-row {{cursor:pointer;user-select:none;}}
  .expandable-row:hover {{background:#eef2ff;}}
  .toggle-icon {{width:20px;text-align:center;font-size:10px;color:var(--accent);transition:transform 0.2s;display:inline-block;}}
  .toggle-icon.open {{transform:rotate(90deg);}}
  .sub-row.hidden {{display:none;}}
  .sub-key {{color:var(--muted);font-size:12px;}}
  .sub-num {{color:var(--muted);font-size:12px;}}
  .sub-badge {{font-size:10px;background:var(--border);color:var(--muted);padding:1px 6px;border-radius:99px;margin-left:6px;font-family:-apple-system,sans-serif;}}
  .footer {{text-align:center;color:var(--muted);font-size:12px;padding:20px 0;}}
</style>
</head>
<body>
<div class="container">
<h1>MaaS 消费明细分析报告</h1>
<p class="subtitle">
  文件: {csv_filename} &nbsp;|&nbsp;
  生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} &nbsp;|&nbsp;
  总计 {total_records:,} 条记录
</p>

<!-- 概览卡片 -->
<div class="cards">
  <div class="card"><div class="label">总记录数</div><div class="value">{total_records:,}</div></div>
  <div class="card"><div class="label">Payer 用户</div><div class="value accent">{unique_users}</div></div>
  <div class="card"><div class="label">Resource ID</div><div class="value accent">{unique_resources}</div></div>
  <div class="card"><div class="label">Instance ID</div><div class="value accent">{unique_instances}</div></div>
  <div class="card"><div class="label">Product</div><div class="value accent">{unique_products}</div></div>
  <div class="card"><div class="label">Project</div><div class="value accent">{unique_projects}</div></div>
  <div class="card"><div class="label">总 Usage ({main_unit})</div><div class="value accent2">{total_usage:,.0f}</div></div>
  <div class="card"><div class="label">平均每条 Usage</div><div class="value accent2">{avg_usage:,.2f}</div></div>
  <div class="card"><div class="label">总税前金额 ({main_currency})</div><div class="value accent3">{total_pretax:,.2f}</div></div>
  <div class="card"><div class="label">总税金 ({main_currency})</div><div class="value accent3">{total_tax:,.2f}</div></div>
  <div class="card"><div class="label">总税后金额 ({main_currency})</div><div class="value accent3">{total_posttax:,.2f}</div></div>
  <div class="card"><div class="label">总应付金额 ({main_currency})</div><div class="value accent3">{total_amount_due:,.2f}</div></div>
  <div class="card"><div class="label">总已付金额 ({main_currency})</div><div class="value accent3">{total_amount_paid:,.2f}</div></div>
  <div class="card card-wide"><div class="label">数据时间范围</div><div class="value time-val">{time_range}</div></div>
</div>

<!-- 柱状图 -->
<div class="charts">
  <div class="chart-box"><h3>Top 10 Payer - 总应付金额</h3><canvas id="chartPayer"></canvas></div>
  <div class="chart-box"><h3>Top 10 Resource ID - 总 Usage</h3><canvas id="chartResource"></canvas></div>
  <div class="chart-box"><h3>Top 10 Product - 总应付金额</h3><canvas id="chartProduct"></canvas></div>
</div>

<!-- 表格 1: 按 Payer User Name -->
<div class="section">
  <h2>按 Payer User Name <span class="badge">{len(stats['payer_user_name'])}</span></h2>
  <div class="table-wrap">
    <table>
      <thead><tr>
        <th>#</th><th>Payer User Name</th><th class="num">记录数</th><th class="num">总 Usage</th><th class="num">平均 Usage</th><th class="num">税前金额</th><th class="num">税</th><th class="num">税后金额</th><th class="num">应付金额</th>
      </tr></thead>
      <tbody>{table_rows(stats['payer_user_name'])}</tbody>
    </table>
  </div>
</div>

<!-- 表格 2: 按 Resource ID -->
<div class="section">
  <h2>按 Resource ID <span class="badge">{len(stats['resource_id'])}</span></h2>
  <div class="table-wrap">
    <table>
      <thead><tr>
        <th>#</th><th>Resource ID</th><th class="num">记录数</th><th class="num">总 Usage</th><th class="num">平均 Usage</th><th class="num">税前金额</th><th class="num">税</th><th class="num">税后金额</th><th class="num">应付金额</th>
      </tr></thead>
      <tbody>{table_rows(stats['resource_id'])}</tbody>
    </table>
  </div>
</div>

<!-- 表格 3: 按 User (user-username) -->
<div class="section">
  <h2>按 User (user-username) <span class="badge">{len(stats['user_name'])}</span></h2>
  <div class="table-wrap">
    <table id="userTable">
      <thead><tr>
        <th></th><th>#</th><th>User</th><th class="num">记录数</th><th class="num">总 Usage</th><th class="num">平均 Usage</th><th class="num">税前金额</th><th class="num">税</th><th class="num">税后金额</th><th class="num">应付金额</th>
      </tr></thead>
      <tbody>{user_table_rows(stats['user_name'])}</tbody>
    </table>
  </div>
</div>

<!-- 表格 4: 按 Product -->
<div class="section">
  <h2>按 Product <span class="badge">{len(stats['product'])}</span></h2>
  <div class="table-wrap">
    <table>
      <thead><tr>
        <th>#</th><th>Product</th><th class="num">记录数</th><th class="num">总 Usage</th><th class="num">平均 Usage</th><th class="num">税前金额</th><th class="num">税</th><th class="num">税后金额</th><th class="num">应付金额</th>
      </tr></thead>
      <tbody>{table_rows(stats['product'])}</tbody>
    </table>
  </div>
</div>

<!-- 表格 5: 按 Instance ID -->
<div class="section">
  <h2>按 Instance ID <span class="badge">{len(stats['instance_id'])}</span></h2>
  <div class="table-wrap">
    <table>
      <thead><tr>
        <th>#</th><th>Instance ID</th><th class="num">记录数</th><th class="num">总 Usage</th><th class="num">平均 Usage</th><th class="num">税前金额</th><th class="num">税</th><th class="num">税后金额</th><th class="num">应付金额</th>
      </tr></thead>
      <tbody>{table_rows(stats['instance_id'])}</tbody>
    </table>
  </div>
</div>

<div class="footer">Generated by parse_and_report.py &mdash; {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</div>
</div>

<!-- Canvas 柱状图 -->
<script>
// ── 可展开行 toggle ──
function toggleSub(rowId, el) {{
  const icon = el.querySelector('.toggle-icon');
  const subRows = document.querySelectorAll('.sub-row[data-parent="' + rowId + '"]');
  const isOpen = !subRows[0].classList.contains('hidden');
  subRows.forEach(r => r.classList.toggle('hidden', isOpen));
  icon.classList.toggle('open', !isOpen);
  icon.innerHTML = isOpen ? '&#9654;' : '&#9660;';
}}

(function() {{
  const charts = [
    {payer_chart},
    {resource_chart},
    {product_chart},
  ];

  const COLORS = ['#4f6ef7','#22c55e','#f59e0b','#ef4444','#8b5cf6','#06b6d4','#ec4899','#14b8a6','#f97316','#6366f1'];

  charts.forEach(c => {{
    const canvas = document.getElementById(c.id);
    if (!canvas) return;
    const ctx = canvas.getContext('2d');
    const dpr = window.devicePixelRatio || 1;
    const rect = canvas.parentElement.getBoundingClientRect();
    canvas.width = rect.width * dpr;
    canvas.height = 300 * dpr;
    canvas.style.width = rect.width + 'px';
    canvas.style.height = '300px';
    ctx.scale(dpr, dpr);

    const w = rect.width, h = 300;
    const pad = {{ top: 30, right: 20, bottom: 80, left: 60 }};
    const cw = w - pad.left - pad.right;
    const ch = h - pad.top - pad.bottom;

    const maxVal = Math.max(...c.values, 1);
    const barW = Math.max(6, cw / c.labels.length * 0.6);
    const gap = cw / c.labels.length;

    // 背景
    ctx.fillStyle = '#fff';
    ctx.fillRect(0, 0, w, h);

    // Y 轴刻度线
    ctx.strokeStyle = '#e5e7eb';
    ctx.lineWidth = 1;
    const ySteps = 5;
    for (let i = 0; i <= ySteps; i++) {{
      const y = pad.top + (ch / ySteps) * i;
      ctx.beginPath();
      ctx.moveTo(pad.left, y);
      ctx.lineTo(w - pad.right, y);
      ctx.stroke();
      ctx.fillStyle = '#888';
      ctx.font = '11px -apple-system,sans-serif';
      ctx.textAlign = 'right';
      const val = (maxVal / ySteps) * (ySteps - i);
      ctx.fillText(val.toFixed(val < 10 ? 2 : 0), pad.left - 6, y + 4);
    }}

    // 柱状条
    c.values.forEach((v, i) => {{
      const x = pad.left + gap * i + (gap - barW) / 2;
      const barH = (v / maxVal) * ch;
      const y = pad.top + ch - barH;
      ctx.fillStyle = COLORS[i % COLORS.length];
      ctx.beginPath();
      ctx.roundRect(x, y, barW, barH, [4, 4, 0, 0]);
      ctx.fill();

      // 值标注
      ctx.fillStyle = '#333';
      ctx.font = 'bold 10px -apple-system,sans-serif';
      ctx.textAlign = 'center';
      ctx.fillText(v.toFixed(v < 10 ? 2 : 0), x + barW / 2, y - 6);
    }});

    // X 轴标签
    ctx.fillStyle = '#555';
    ctx.font = '10px -apple-system,sans-serif';
    c.labels.forEach((label, i) => {{
      const x = pad.left + gap * i + gap / 2;
      ctx.save();
      ctx.translate(x, pad.top + ch + 16);
      ctx.rotate(-0.5);
      const txt = label.length > 28 ? label.slice(0, 26) + '..' : label;
      ctx.fillText(txt, 0, 0);
      ctx.restore();
    }});
  }});
}})();
</script>
</body>
</html>"""
    return html


# ── 入口 ──────────────────────────────────────────────────────────────
def main():
    if len(sys.argv) < 2:
        print("用法: python3 parse_and_report.py <CSV文件路径>")
        sys.exit(1)

    csv_path = sys.argv[1]
    if not os.path.exists(csv_path):
        print(f"[ERROR] 找不到文件: {csv_path}")
        sys.exit(1)

    # 输出文件名: 输入文件名_时间戳.html
    base_name = os.path.splitext(os.path.basename(csv_path))[0]
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    script_dir = os.path.dirname(os.path.abspath(__file__))
    out_path = os.path.join(script_dir, f"{base_name}_{timestamp}.html")

    # 加载 Resource ID 名称映射
    mapping, resource_to_username = load_mapping("name-apikey-mapping.csv")

    print(f"正在解析: {csv_path}")
    rows = parse_csv(csv_path, resource_to_username)
    print(f"共解析 {len(rows)} 条记录")

    # 多维度聚合
    dimensions = ["payer_user_name", "resource_id", "user_name", "product", "instance_id"]
    stats = {}
    for dim in dimensions:
        print(f"  聚合维度: {dim}")
        stats[dim] = aggregate(rows, dim, top_n=30, mapping=mapping)

    # 生成 HTML
    print(f"正在生成 HTML: {out_path}")
    html = build_html(rows, stats, os.path.basename(csv_path))
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)

    print(f"[DONE] 报告已生成: {out_path}")
    print(f"       大小: {os.path.getsize(out_path):,} bytes")


if __name__ == "__main__":
    main()
